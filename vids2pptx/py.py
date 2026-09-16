#!/usr/bin/env python3
"""
video2pptx.py - turn a recorded presentation (MP4) into a PowerPoint deck.

Two steps, two commands:

    analyze   MP4 in GCS  -->  Gemini on Vertex AI  -->  slides.json
    build     slides.json + local MP4  -->  ffmpeg frames  -->  deck.pptx

They are deliberately separate. `analyze` costs money and takes minutes;
`build` is free and takes seconds. So you can hand-edit slides.json (fix a
timestamp, reword a bullet, delete a slide) and rebuild as often as you like
without calling the model again.

How it works
------------
Gemini ingests the whole video, audio included, sampled at one frame per
second with a timestamp on every second. It is asked for a JSON deck: per
slide, a title, terse bullets, the full spoken instruction as speaker notes,
and - when something on screen matters - the MM:SS of the frame that best
shows it. That timestamp is the join key between the model's understanding
and the pixels: `build` seeks to it with ffmpeg, grabs one frame, and lays
the slide out with python-pptx. Nothing has to understand the video twice.

Usage
-----
    python video2pptx.py analyze --video talk.mp4 --project MY_PROJECT --bucket MY_BUCKET
    python video2pptx.py build   --video talk.mp4 --json slides.json --out talk.pptx

Run `python video2pptx.py analyze --help` / `build --help` for all flags.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

# --------------------------------------------------------------------------
# Configuration you may want to tweak
# --------------------------------------------------------------------------

DEFAULT_MODEL = "gemini-3.5-flash"   # any current Gemini 3 model; Pro for higher quality
DEFAULT_LOCATION = "global"          # or a region like "us-central1" if policy requires
DEFAULT_MAX_SLIDES = 40
MAX_GAP_MINUTES = 3                  # consecutive slides further apart than this = the model skipped something
TOKENS_PER_SECOND_DEFAULT_RES = 100  # rough Gemini 3 cost of video at default resolution, for sanity checks
GCS_PREFIX = "video2pptx"            # folder inside the bucket for uploads
DEFAULT_RUNS_DIR = "runs"            # each analyze creates runs/<video>_<timestamp>/ holding everything for that run
REQUEST_TIMEOUT_MS = 30 * 60 * 1000  # a 50-minute video can take several minutes to process

STYLE = {
    "font": "Calibri",
    "title_pt": 30,
    "body_pt": 18,
    "caption_pt": 11,
    "meta_pt": 9,
    "text_rgb": (0x1F, 0x1F, 0x1F),
    "muted_rgb": (0x6B, 0x6B, 0x6B),
}

PROMPT_TEMPLATE = """\
You are converting a recorded instructional presentation into a PowerPoint deck.
The video is about {minutes} minutes long.{context}

Rules:
- Work through the whole video in order, start to finish. Do not skip sections.
- One slide per distinct step, concept, or decision the presenter explains. Produce between
  {min_slides} and {max_slides} slides. The most common mistake is producing too few: do NOT
  merge several steps into one slide, and do NOT simply mirror the slides the presenter shows
  on screen - segment by what is SAID. Roughly one slide per 1-2 minutes of speech.
- starts_at: the timestamp (MM:SS, or H:MM:SS past 59 minutes) where this slide's content
  begins in the video. Slides must be in increasing order of starts_at, and no two consecutive
  slides should be more than {max_gap} minutes apart - if they are, you have skipped something.
- title: short and specific (max 8 words).
- bullets: 3-5 terse lines (max 12 words each) capturing what to do or know. No sub-bullets.
- notes: the FULL spoken instruction for this slide, rewritten as clear speaker notes in
  second person. Preserve details, exact names, menu paths, values, and warnings.
  This is the primary record of what was said - be complete, not brief.
- frame_at: if the presenter shows something on screen that matters for this slide
  (a screen, dialog, diagram, demo), give the timestamp of the frame that best shows it,
  a couple of seconds AFTER the screen has settled. Use the absolute position in the video
  as MM:SS (for example 07:45), or H:MM:SS past 59 minutes. If nothing visual matters for
  this slide, use an empty string.
- frame_caption: one short line saying what the frame shows, or an empty string.
- deck_title: a short title for the whole deck.
"""

RESPONSE_SCHEMA = {
    "type": "object",
    "properties": {
        "deck_title": {"type": "string"},
        "slides": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "starts_at": {"type": "string"},
                    "title": {"type": "string"},
                    "bullets": {"type": "array", "items": {"type": "string"}},
                    "notes": {"type": "string"},
                    "frame_at": {"type": "string"},
                    "frame_caption": {"type": "string"},
                },
                "required": ["starts_at", "title", "bullets", "notes", "frame_at", "frame_caption"],
            },
        },
    },
    "required": ["deck_title", "slides"],
}

MEDIA_RESOLUTION = {  # CLI value -> google.genai.types.MediaResolution member name
    "low": "MEDIA_RESOLUTION_LOW",
    "medium": "MEDIA_RESOLUTION_MEDIUM",
    "high": "MEDIA_RESOLUTION_HIGH",
}


# --------------------------------------------------------------------------
# Small shared helpers
# --------------------------------------------------------------------------

_LOG_FILE: Path | None = None   # set by start_run_log(); everything logged is also appended there


def log(msg: str) -> None:
    print(msg, flush=True)
    if _LOG_FILE:
        with _LOG_FILE.open("a", encoding="utf-8") as f:
            f.write(msg + "\n")


def die(msg: str, code: int = 1) -> None:
    print(f"ERROR: {msg}", file=sys.stderr, flush=True)
    if _LOG_FILE:
        with _LOG_FILE.open("a", encoding="utf-8") as f:
            f.write(f"ERROR: {msg}\n")
    sys.exit(code)


def start_run_log(run_dir: Path, command: str, argv: list[str]) -> None:
    """Append a dated header to <run>/run.log and tee all further log() output into it,
    so the coverage warnings and build messages survive after the terminal scrolls."""
    global _LOG_FILE
    _LOG_FILE = run_dir / "run.log"
    with _LOG_FILE.open("a", encoding="utf-8") as f:
        f.write(f"\n=== {command}  {dt.datetime.now():%Y-%m-%d %H:%M:%S}  {' '.join(argv)}\n")


def new_run_dir(runs_dir: Path, video: Path, name: str | None) -> Path:
    """runs/<name>/ or runs/<video-stem>_<YYYYMMDD-HHMMSS>/, created empty."""
    run_dir = runs_dir / (name or f"{video.stem}_{dt.datetime.now():%Y%m%d-%H%M%S}")
    if run_dir.exists() and any(run_dir.iterdir()):
        die(f"Run folder already exists and is not empty: {run_dir}. Pick another --run name.")
    run_dir.mkdir(parents=True, exist_ok=True)
    return run_dir


def latest_run_dir(runs_dir: Path) -> Path:
    """The most recently modified run folder that contains a slides.json."""
    candidates = [d for d in runs_dir.iterdir() if (d / "slides.json").is_file()] if runs_dir.is_dir() else []
    if not candidates:
        die(f"No runs found under {runs_dir}/. Run `analyze` first, or pass --run / --json.")
    return max(candidates, key=lambda d: (d / "slides.json").stat().st_mtime)


def require_tool(name: str) -> None:
    if shutil.which(name) is None:
        die(f"'{name}' is not on your PATH. Install ffmpeg (which includes ffprobe) and open a new terminal.")


def video_duration_seconds(video: Path) -> float | None:
    """Ask ffprobe for the duration. Returns None if ffprobe is unavailable."""
    if shutil.which("ffprobe") is None:
        return None
    out = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration",
         "-of", "default=noprint_wrappers=1:nokey=1", str(video)],
        capture_output=True, text=True,
    )
    try:
        return float(out.stdout.strip())
    except ValueError:
        return None


_TS_RE = re.compile(r"^\s*(?:(\d{1,2}):)?(\d{1,3}):(\d{2})(?:\.(\d+))?\s*$")


def parse_timestamp(value: str | None) -> float | None:
    """'07:45' / '1:02:30' / '07:45.5' -> seconds. Anything else -> None."""
    if not value:
        return None
    m = _TS_RE.match(value)
    if not m:
        return None
    hours = int(m.group(1) or 0)
    minutes = int(m.group(2))
    seconds = int(m.group(3))
    frac = float("0." + m.group(4)) if m.group(4) else 0.0
    return hours * 3600 + minutes * 60 + seconds + frac


def fmt_timestamp(seconds: float) -> str:
    seconds = max(0.0, seconds)
    h, rem = divmod(int(seconds), 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


# --------------------------------------------------------------------------
# ANALYZE: video --> Gemini (Vertex AI) --> slides.json
# --------------------------------------------------------------------------

def upload_to_gcs(video: Path, bucket_name: str, project: str) -> str:
    """Upload the MP4 to gs://bucket/video2pptx/<name>, skipping if already there."""
    from google.cloud import storage  # imported here so `build` never needs it

    client = storage.Client(project=project)
    blob = client.bucket(bucket_name).blob(f"{GCS_PREFIX}/{video.name}")
    uri = f"gs://{bucket_name}/{blob.name}"

    local_size = video.stat().st_size
    if blob.exists():
        blob.reload()
        if blob.size == local_size:
            log(f"Already in GCS (same size), skipping upload: {uri}")
            return uri
        log("Object exists but size differs; re-uploading.")

    log(f"Uploading {video.name} ({local_size / 1e6:.0f} MB) to {uri} ...")
    blob.chunk_size = 32 * 1024 * 1024  # resumable upload in 32 MB chunks
    blob.upload_from_filename(str(video), timeout=600)
    log("Upload complete.")
    return uri


def cmd_analyze(args: argparse.Namespace) -> None:
    try:
        from google import genai
        from google.genai import types
    except ImportError:
        die("google-genai is not installed. Run: pip install -r requirements.txt")

    video = Path(args.video).resolve()
    if not video.is_file():
        die(f"Video not found: {video}")

    run_dir = new_run_dir(Path(args.runs_dir), video, args.run)
    start_run_log(run_dir, "analyze", sys.argv[1:])
    log(f"Run folder: {run_dir}")

    project = args.project or os.environ.get("GOOGLE_CLOUD_PROJECT")
    if not project:
        die("Pass --project or set GOOGLE_CLOUD_PROJECT.")
    location = args.location or os.environ.get("GOOGLE_CLOUD_LOCATION") or DEFAULT_LOCATION

    # 1. Get the video somewhere Vertex can read it.
    if args.gcs_uri:
        gcs_uri = args.gcs_uri
    elif args.bucket:
        gcs_uri = upload_to_gcs(video, args.bucket, project)
    else:
        die("Pass --bucket (script uploads for you) or --gcs-uri gs://... (already uploaded).")

    # 2. Build the prompt.
    duration = video_duration_seconds(video)
    minutes = f"{duration / 60:.0f}" if duration else "unknown"
    context = f"\nContext from the requester: {args.context.strip()}" if args.context else ""
    if args.min_slides is not None:
        min_slides = args.min_slides
    elif duration:
        min_slides = max(5, round(duration / 60 / 2))      # one per 2 minutes is the floor
    else:
        min_slides = 5
    min_slides = min(min_slides, args.max_slides)
    prompt = PROMPT_TEMPLATE.format(minutes=minutes, context=context, min_slides=min_slides,
                                    max_slides=args.max_slides, max_gap=MAX_GAP_MINUTES)
    (run_dir / "prompt.txt").write_text(prompt, encoding="utf-8")   # exactly what the model was asked

    # 3. Call Gemini.
    client = genai.Client(
        vertexai=True, project=project, location=location,
        http_options=types.HttpOptions(timeout=REQUEST_TIMEOUT_MS),
    )
    config_kwargs = dict(
        response_mime_type="application/json",
        response_schema=RESPONSE_SCHEMA,
        temperature=0.2,             # we want faithful extraction, not creativity
        max_output_tokens=65535,     # 40 slides of full speaker notes is a lot of text
        # We pass no tools, so AFC is irrelevant; saying so explicitly silences the SDK's warning.
        automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True),
    )
    if args.resolution:
        config_kwargs["media_resolution"] = getattr(types.MediaResolution, MEDIA_RESOLUTION[args.resolution])

    log(f"Calling {args.model} in {location} on {gcs_uri} "
        f"(~{minutes} min of video; this typically takes a few minutes) ...")
    started = dt.datetime.now()
    response = client.models.generate_content(
        model=args.model,
        contents=[
            types.Part.from_uri(file_uri=gcs_uri, mime_type="video/mp4"),
            prompt,
        ],
        config=types.GenerateContentConfig(**config_kwargs),
    )
    elapsed = (dt.datetime.now() - started).total_seconds()

    # 4. Parse and save.
    raw = response.text or ""
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as e:
        raw_path = run_dir / "slides.raw.txt"
        raw_path.write_text(raw, encoding="utf-8")
        die(f"Model output was not valid JSON ({e}). Raw text saved to {raw_path}. "
            f"If it looks truncated, lower --max-slides and retry.")

    usage = getattr(response, "usage_metadata", None)
    data["_meta"] = {
        "source_video": video.name,
        "source_video_path": str(video),      # lets `build --run X` find the file without --video
        "gcs_uri": gcs_uri,
        "model": args.model,
        "location": location,
        "media_resolution": args.resolution or "default",
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "elapsed_seconds": round(elapsed),
        "prompt_tokens": getattr(usage, "prompt_token_count", None),
        "output_tokens": getattr(usage, "candidates_token_count", None),
    }
    json_path = run_dir / "slides.json"
    json_path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

    slides = data.get("slides", [])
    with_frames = sum(1 for s in slides if parse_timestamp(s.get("frame_at")) is not None)
    log(f"Done in {elapsed:.0f}s. {len(slides)} slides ({with_frames} with a frame) -> {json_path}")
    if usage:
        log(f"Tokens: {data['_meta']['prompt_tokens']} in / {data['_meta']['output_tokens']} out")
    report_coverage(slides, duration, min_slides, data["_meta"].get("prompt_tokens"))
    log(f"Next: open {json_path}, sanity-check a few slides, then run `build` (it will pick this run by default).")


def report_coverage(slides: list, duration: float | None, min_slides: int, prompt_tokens: int | None) -> None:
    """Print where the slides fall on the video's timeline and flag the three ways a run
    goes wrong: the model didn't ingest the whole video, it stopped early, or it merged
    too much into too few slides. Each check is cheap and points at a different fix."""
    starts = [parse_timestamp(s.get("starts_at")) for s in slides]
    starts = [t for t in starts if t is not None]
    problems = []

    if duration and prompt_tokens:
        expected = duration * TOKENS_PER_SECOND_DEFAULT_RES
        if prompt_tokens < 0.3 * expected:
            problems.append(f"prompt used {prompt_tokens:,} tokens but ~{expected:,.0f} were expected for "
                            f"{duration/60:.0f} min of video - the model may not have ingested the whole file "
                            f"(check that --video is the full recording and that it has an audio stream: ffprobe FILE)")

    if starts:
        first, last = min(starts), max(starts)
        gaps = [(b - a, i) for i, (a, b) in enumerate(zip(starts, starts[1:]), start=1)]
        worst_gap, at = max(gaps) if gaps else (0.0, 0)
        where = f" of {fmt_timestamp(duration)}" if duration else ""
        log(f"Timeline: slides span {fmt_timestamp(first)} - {fmt_timestamp(last)}{where}; "
            f"largest gap {fmt_timestamp(worst_gap)} (between slides {at} and {at + 1})")
        if duration and last < 0.7 * duration:
            problems.append(f"the last slide starts at {fmt_timestamp(last)} of {fmt_timestamp(duration)} - "
                            f"the model stopped early or only saw part of the video")
        if worst_gap > MAX_GAP_MINUTES * 60 * 1.5:
            problems.append(f"a {fmt_timestamp(worst_gap)} stretch has no slide - content was probably merged or skipped")

    if len(slides) < min_slides:
        problems.append(f"only {len(slides)} slides; the prompt asked for at least {min_slides}. "
                        f"Re-run with --context describing the steps you expect, or a higher --min-slides")

    for p in problems:
        log(f"  WARNING: {p}")


# --------------------------------------------------------------------------
# BUILD: slides.json + local MP4 --> frames --> deck.pptx
# --------------------------------------------------------------------------

def extract_frame(video: Path, seconds: float, out: Path) -> None:
    """Grab one JPEG at `seconds`. `-ss` before `-i` is fast (keyframe seek) AND
    frame-accurate in modern ffmpeg, because it decodes forward from the keyframe
    and discards frames until the exact time. -q:v 2 = high-quality JPEG."""
    subprocess.run(
        ["ffmpeg", "-y", "-loglevel", "error",
         "-ss", f"{seconds:.3f}", "-i", str(video),
         "-frames:v", "1", "-q:v", "2", str(out)],
        check=True,
    )


def cmd_build(args: argparse.Namespace) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    require_tool("ffmpeg")

    # Which run? An explicit --json wins; else --run; else the newest run folder.
    if args.json:
        json_path = Path(args.json)
        run_dir = json_path.parent
    else:
        run_dir = Path(args.run) if args.run else latest_run_dir(Path(args.runs_dir))
        json_path = run_dir / "slides.json"
    if not json_path.is_file():
        die(f"JSON not found: {json_path} (run `analyze` first)")
    start_run_log(run_dir, "build", sys.argv[1:])
    log(f"Run folder: {run_dir}")

    data = json.loads(json_path.read_text(encoding="utf-8"))
    slides = data.get("slides") or []
    if not slides:
        die("No slides in JSON.")
    meta = data.get("_meta") or {}

    # The video: --video if given, else the path analyze recorded for this run.
    video = Path(args.video or meta.get("source_video_path") or "")
    if not str(video) or not video.is_file():
        die(f"Video not found ({video or 'no path recorded in the run'}). Pass --video.")
    deck_title = data.get("deck_title") or video.stem
    duration = video_duration_seconds(video)

    out = Path(args.out) if args.out else run_dir / f"{video.stem}.pptx"
    frames_dir = Path(args.frames_dir) if args.frames_dir else run_dir / "frames"
    frames_dir.mkdir(parents=True, exist_ok=True)

    # ---- layout constants (16:9, 13.333 x 7.5 in) ----
    W, H = Inches(13.333), Inches(7.5)
    M = Inches(0.6)                      # outer margin
    TITLE_TOP, TITLE_H = Inches(0.45), Inches(1.0)
    BODY_TOP = Inches(1.65)
    BODY_H = Inches(4.9)
    TEXT_W_WITH_PIC = Inches(5.6)
    PIC_LEFT = Inches(6.7)
    PIC_W = W - PIC_LEFT - M
    PIC_H = Inches(4.3)
    TEXT_RGB = RGBColor(*STYLE["text_rgb"])
    MUTED_RGB = RGBColor(*STYLE["muted_rgb"])

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    def textbox(slide, left, top, width, height, text, pt, *, bold=False,
                color=TEXT_RGB, align=None, anchor=None, italic=False):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        if anchor is not None:
            tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        if align is not None:
            p.alignment = align
        r = p.add_run()             # a run, not tf.text, so formatting sticks
        r.text = text
        r.font.name = STYLE["font"]
        r.font.size = Pt(pt)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        return tf

    def bullets(slide, left, top, width, height, items):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            # A real bullet (a:buChar) with a hanging indent - not a literal "•" in the text,
            # which would render as a double bullet if the layout ever adds its own.
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.3)))
            pPr.set("indent", str(-Inches(0.3)))
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u2022"}))
            r = p.add_run()
            r.text = item
            r.font.name = STYLE["font"]
            r.font.size = Pt(STYLE["body_pt"])
            r.font.color.rgb = TEXT_RGB

    def picture_fit(slide, path, left, top, max_w, max_h):
        """Add a picture scaled to fit inside the box, preserving aspect ratio, centred."""
        pic = slide.shapes.add_picture(str(path), left, top, width=max_w)
        if pic.height > max_h:
            scale = max_h / pic.height
            pic.height = int(max_h)
            pic.width = int(pic.width * scale)
        pic.left = int(left + (max_w - pic.width) / 2)
        return pic

    # ---- title slide ----
    s = prs.slides.add_slide(blank)
    textbox(s, M, Inches(2.3), W - 2 * M, Inches(1.6), deck_title, 40,
            bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, M, Inches(4.1), W - 2 * M, Inches(0.6),
            f"Generated from {video.name} \u00b7 {dt.date.today():%B %d, %Y}",
            16, color=MUTED_RGB, align=PP_ALIGN.CENTER)

    # ---- content slides ----
    total = len(slides)
    frames_used = 0
    for n, sd in enumerate(slides, start=1):
        s = prs.slides.add_slide(blank)
        title = (sd.get("title") or f"Slide {n}").strip()
        items = [b.strip() for b in (sd.get("bullets") or []) if b and b.strip()]
        notes = (sd.get("notes") or "").strip()
        caption = (sd.get("frame_caption") or "").strip()

        secs = parse_timestamp(sd.get("frame_at"))
        if secs is not None and duration:
            secs = min(secs, max(0.0, duration - 0.5))   # clamp to inside the video

        textbox(s, M, TITLE_TOP, W - 2 * M, TITLE_H, title, STYLE["title_pt"],
                bold=True, anchor=MSO_ANCHOR.MIDDLE)

        frame_path = None
        if secs is not None:
            frame_path = frames_dir / f"slide{n:02d}.jpg"
            try:
                extract_frame(video, secs, frame_path)
                frames_used += 1
            except subprocess.CalledProcessError:
                log(f"  warning: could not extract frame for slide {n} at {fmt_timestamp(secs)}")
                frame_path = None

        if frame_path:
            bullets(s, M, BODY_TOP, TEXT_W_WITH_PIC, BODY_H, items)
            pic = picture_fit(s, frame_path, PIC_LEFT, BODY_TOP, PIC_W, PIC_H)
            if caption:
                textbox(s, PIC_LEFT, pic.top + pic.height + Inches(0.1), PIC_W, Inches(0.5),
                        caption, STYLE["caption_pt"], color=MUTED_RGB, italic=True,
                        align=PP_ALIGN.CENTER)
        else:
            bullets(s, M, BODY_TOP, W - 2 * M, BODY_H, items)

        # footer: slide number (left) and where in the video this came from (right)
        textbox(s, M, H - Inches(0.5), Inches(3), Inches(0.3), f"{n} / {total}",
                STYLE["meta_pt"], color=MUTED_RGB)
        if secs is not None:
            textbox(s, W - M - Inches(3), H - Inches(0.5), Inches(3), Inches(0.3),
                    f"Source video @ {fmt_timestamp(secs)}", STYLE["meta_pt"],
                    color=MUTED_RGB, align=PP_ALIGN.RIGHT)

        if notes:
            s.notes_slide.notes_text_frame.text = notes

        log(f"  slide {n:02d}/{total}: {title}" + (f"  [frame @ {fmt_timestamp(secs)}]" if frame_path else ""))

    prs.save(str(out))
    log(f"Saved {out} ({total} content slides, {frames_used} frames). Frames kept in {frames_dir}/")
    log(f"Next: open {out}. To fix a slide, edit {json_path} (timestamp, bullets, notes) and run `build` again -")
    log("      it is free and takes seconds. Re-run `analyze` only if the overall shape of the deck is wrong.")


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

QUICKSTART = """\
QUICK START  (PowerShell, from the folder containing this script)

  1. Activate the environment           .\\.venv\\Scripts\\Activate.ps1
  2. Try a 5-minute clip first          ffmpeg -i talk.mp4 -t 300 -c copy test5.mp4
  3. Analyze  (video -> slides.json)    python video2pptx.py analyze --video test5.mp4 --project MY_PROJECT --bucket MY_BUCKET
  4. Build    (json  -> deck.pptx)      python video2pptx.py build
  5. Happy? Repeat 3-4 on the full video.  Not happy? Edit the run's slides.json and re-run step 4.

  Every analyze creates runs/<video>_<timestamp>/ holding prompt.txt, slides.json, run.log,
  frames/ and the built deck. `build` uses the newest run unless you pass --run <folder>.

  python video2pptx.py analyze --help     all analyze options (context, min/max slides, resolution ...)
  python video2pptx.py build --help       all build options
  See RUN_GUIDE.md for setup, tuning, and troubleshooting.
"""

ANALYZE_EXAMPLES = """\
examples:
  # simplest: script uploads the video for you; output lands in runs/talk_<timestamp>/
  python video2pptx.py analyze --video talk.mp4 --project MY_PROJECT --bucket MY_BUCKET

  # name the run yourself
  python video2pptx.py analyze --video talk.mp4 --project MY_PROJECT --bucket MY_BUCKET --run talk-highres

  # with context (recommended) and a slide ceiling
  python video2pptx.py analyze --video talk.mp4 --project MY_PROJECT --bucket MY_BUCKET ^
      --context "Audience: new hires. Walkthrough of the supplier portal." --max-slides 30

  # video already in GCS; small on-screen text matters
  python video2pptx.py analyze --video talk.mp4 --project MY_PROJECT ^
      --gcs-uri gs://MY_BUCKET/video2pptx/talk.mp4 --resolution high

(^ is PowerShell/cmd line continuation; on one line just omit it.)
"""

BUILD_EXAMPLES = """\
examples:
  python video2pptx.py build                                  # newest run -> runs/<run>/<video>.pptx
  python video2pptx.py build --run runs/talk_20260916-143205  # a specific run
  python video2pptx.py build --run runs/talk_20260916-143205 --out C:\\share\\talk-v2.pptx

Edit the run's slides.json between builds: change a slide's "frame_at" to move its screenshot,
set it to "" to drop the picture, reword "bullets", or delete a slide object entirely.
"""


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(
        description="Turn a recorded presentation into a PowerPoint deck with Gemini on Vertex AI.",
        epilog=QUICKSTART,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    sub = ap.add_subparsers(dest="command")

    a = sub.add_parser("analyze", help="Send the video to Gemini and write slides.json",
                       epilog=ANALYZE_EXAMPLES, formatter_class=argparse.RawDescriptionHelpFormatter)
    a.add_argument("--video", required=True, help="Local MP4 (used for its name and duration)")
    src = a.add_mutually_exclusive_group()
    src.add_argument("--bucket", help="GCS bucket name; the script uploads the video there")
    src.add_argument("--gcs-uri", help="gs://bucket/path.mp4 if you already uploaded it")
    a.add_argument("--project", help="GCP project ID (or set GOOGLE_CLOUD_PROJECT)")
    a.add_argument("--location", help=f"Vertex location (default: {DEFAULT_LOCATION})")
    a.add_argument("--model", default=DEFAULT_MODEL, help=f"Model ID (default: {DEFAULT_MODEL})")
    a.add_argument("--resolution", choices=list(MEDIA_RESOLUTION), default=None,
                   help="Video frame detail: low (cheapest), medium, high (reads small on-screen text). "
                        "Default: let the model choose.")
    a.add_argument("--context", default="",
                   help='Free text for the model, e.g. "Audience: new hires. This is a walkthrough of the X tool."')
    a.add_argument("--min-slides", type=int, default=None,
                   help="Floor on slide count (default: about one per 2 minutes of video, at least 5)")
    a.add_argument("--max-slides", type=int, default=DEFAULT_MAX_SLIDES,
                   help=f"Ceiling on slide count (default: {DEFAULT_MAX_SLIDES})")
    a.add_argument("--run", help="Name for this run's folder (default: <video>_<timestamp>)")
    a.add_argument("--runs-dir", default=DEFAULT_RUNS_DIR, help=f"Parent folder for runs (default: {DEFAULT_RUNS_DIR}/)")
    a.set_defaults(func=cmd_analyze)

    b = sub.add_parser("build", help="Turn slides.json + the local MP4 into a .pptx",
                       epilog=BUILD_EXAMPLES, formatter_class=argparse.RawDescriptionHelpFormatter)
    b.add_argument("--run", help="Run folder to build from (default: the newest under --runs-dir)")
    b.add_argument("--runs-dir", default=DEFAULT_RUNS_DIR, help=f"Parent folder for runs (default: {DEFAULT_RUNS_DIR}/)")
    b.add_argument("--video", help="Local MP4 (default: the path recorded by analyze for this run)")
    b.add_argument("--json", help="Explicit slides.json to build from (its folder becomes the run folder)")
    b.add_argument("--out", help="Output .pptx (default: <run>/<video>.pptx)")
    b.add_argument("--frames-dir", help="Where frames are written (default: <run>/frames)")
    b.set_defaults(func=cmd_build)

    args = ap.parse_args(argv)
    if args.command is None:          # bare `python video2pptx.py` -> show the quick start, not an error
        print(QUICKSTART)
        return
    args.func(args)


if __name__ == "__main__":
    main()
